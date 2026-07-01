import datetime
import os
import re
import shutil
import time
import unicodedata
import urllib.parse
from concurrent.futures import ThreadPoolExecutor, as_completed
from typing import List, Optional, Tuple
from urllib.parse import urljoin, urlparse, urlencode

import pandas as pd
import requests
from bs4 import BeautifulSoup
from dotenv import load_dotenv

from services.auth_service import PGE_BASE_URL

load_dotenv()

PGE_ISSUES_PATH = f"{PGE_BASE_URL}/issues"
PGE_BASE_PATH = urlparse(PGE_BASE_URL).path.rstrip("/")
PGE_CF28_ENUMERATIONS_URL = f"{PGE_BASE_URL}/custom_fields/28/enumerations"

# Tracker IDs excluded from filter (Projeto, Ajuste, Task, Falha, Backlog, Mudança de Escopo)
_EXCLUDED_TRACKERS = ["20", "18", "3", "12", "14", "15"]

_DEFAULT_CF28_OPTIONS = [("ITEM 1 - JUNHO - 2026", "222")]

RELATORIO_PDF_NOME = "Relatório de Atividades Geradas.pdf"
RELATORIO_INDIVIDUAIS_DIR = "Relatório de Atividades Individuais"
EVIDENCIAS_DIR = "Evidências"
TAREFAS_GERAL_NOME = "TAREFAS_GERAL.xlsx"
TAREFAS_PROCESSADAS_NOME = "TAREFAS_PROCESSADAS.csv"

_TAREFAS_COLUMNS = ["ID", "Colaborador Atribuído", "Título", "UST", "Catálogo", "Perfil"]

_app_run = True


def set_app_run(value: bool) -> None:
    global _app_run
    _app_run = value


def _timestamp() -> str:
    return datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def _issue_workers() -> int:
    try:
        n = int(os.getenv("THREADS", "1"))
    except ValueError:
        n = 1
    return max(1, min(n, 8))


def _clone_session(session: requests.Session) -> requests.Session:
    worker = requests.Session()
    worker.cookies.update(session.cookies)
    worker.headers.update(session.headers)
    worker.verify = session.verify
    return worker


def _http_get(session: requests.Session, url: str, **kwargs) -> requests.Response:
    if "timeout" not in kwargs:
        kwargs["timeout"] = 60
    r = session.get(url, **kwargs)
    r.raise_for_status()
    return r


def _safe_mkdir(path: str) -> None:
    os.makedirs(path, exist_ok=True)


def _safe_dirname(name: str) -> str:
    if name is None:
        return "desconhecido"
    name = str(name).strip()
    name = urllib.parse.unquote(name)
    name = unicodedata.normalize("NFKD", name)
    name = "".join([c for c in name if not unicodedata.combining(c)])
    name = re.sub(r"[^a-zA-Z0-9 \-\_\(\)\.\,]", "_", name)
    name = re.sub(r"\s+", " ", name).strip()
    return name[:150] if len(name) > 150 else name


def _valid_filename(filename: str) -> str:
    if not filename:
        return "arquivo"
    name = urllib.parse.unquote(str(filename))
    name = unicodedata.normalize("NFKD", name)
    name = "".join([c for c in name if not unicodedata.combining(c)])
    name = re.sub(r'[<>:"/\\|?*\x00-\x1F]', "_", name)
    name = name.strip().strip(".")
    return name[-150:] if name else "arquivo"


def _parse_decimal(value: str) -> str:
    """Normalizes decimal values to use '.' as the separator (e.g. '1,5' -> '1.5')."""
    if not value:
        return ""
    s = str(value).strip().replace("\xa0", "").replace(" ", "")
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".")
    elif "," in s:
        s = s.replace(",", ".")
    return s


def _download_file(session: requests.Session, url: str, dest_path: str, max_retries: int = 3) -> None:
    dest_dir = os.path.dirname(dest_path)
    _safe_mkdir(dest_dir)
    last_exc = None

    for attempt in range(1, max_retries + 1):
        try:
            with session.get(url, stream=True, timeout=120, allow_redirects=True) as r:
                r.raise_for_status()
                with open(dest_path, "wb") as f:
                    for chunk in r.iter_content(chunk_size=512 * 1024):
                        if chunk:
                            f.write(chunk)
            print(f"[{_timestamp()}] OK: {url} -> {dest_path}")
            return
        except Exception as e:
            last_exc = e
            print(f"[{_timestamp()}] WARN: download attempt {attempt}/{max_retries} failed for {url}: {e}")
            time.sleep(1.5 * attempt)
    print(f"[{_timestamp()}] ERROR: failed to download {url} to {dest_path}: {last_exc}")


def build_pge_filter_url(cf_28_id: str) -> str:
    """Builds the PGE issues filter URL; only cf_28 value varies per month."""
    params = [
        ("set_filter", "1"),
        ("sort", "id:desc"),
        ("f[]", "tracker_id"),
        ("op[tracker_id]", "!"),
    ]
    for tracker_id in _EXCLUDED_TRACKERS:
        params.append(("v[tracker_id][]", tracker_id))
    params.extend([
        ("f[]", "cf_28"),
        ("op[cf_28]", "="),
        ("v[cf_28][]", cf_28_id),
        ("f[]", "status_id"),
        ("op[status_id]", "c"),
        ("f[]", ""),
        ("c[]", "parent"),
        ("c[]", "project"),
        ("c[]", "tracker"),
        ("c[]", "status"),
        ("c[]", "cf_23"),
        ("c[]", "closed_on"),
        ("c[]", "cf_20"),
        ("c[]", "cf_28"),
        ("c[]", "cf_22"),
        ("group_by", ""),
        ("t[]", "estimated_hours"),
        ("t[]", "spent_hours"),
        ("t[]", "cf_17"),
        ("t[]", "cf_21"),
        ("t[]", "cf_22"),
        ("t[]", ""),
    ])
    return f"{PGE_ISSUES_PATH}?{urlencode(params)}"

def build_pge_filter_pdf_url(cf_28_id: str) -> str:
    """Builds the PGE issues filter PDF URL (same params as HTML list)."""
    return build_pge_filter_url(cf_28_id).replace("/issues?", "/issues.pdf?", 1)


def _default_cf28_enumerations() -> List[dict]:
    return [
        {"id": enum_id, "label": label, "active": True}
        for label, enum_id in _DEFAULT_CF28_OPTIONS
    ]


def _parse_cf28_enumerations(soup: BeautifulSoup) -> List[dict]:
    enumerations: List[dict] = []
    name_pattern = re.compile(r"custom_field_enumerations\[(\d+)\]\[name\]")
    active_pattern = re.compile(r"custom_field_enumerations\[(\d+)\]\[active\]")

    for li in soup.select("#custom_field_enumerations li"):
        name_input = li.find("input", attrs={"name": name_pattern})
        if not name_input:
            continue
        match = name_pattern.search(name_input.get("name", ""))
        if not match:
            continue
        enum_id = match.group(1)
        label = (name_input.get("value") or "").strip()
        if not label:
            continue

        active = False
        active_cb = li.find("input", attrs={"type": "checkbox", "name": active_pattern})
        if active_cb and active_cb.has_attr("checked"):
            active = True

        enumerations.append({"id": enum_id, "label": label, "active": active})

    return enumerations


def fetch_cf28_enumerations(session: requests.Session) -> List[dict]:
    """Fetches cf_28 enumeration options from the PGE custom field admin page."""
    try:
        print(f"[{_timestamp()}] Buscando enumerações cf_28: {PGE_CF28_ENUMERATIONS_URL}")
        r = _http_get(session, PGE_CF28_ENUMERATIONS_URL)
        if "/login" in r.url:
            print(f"[{_timestamp()}] WARN: sessão sem acesso à página de enumerações; usando fallback.")
            return _default_cf28_enumerations()

        soup = BeautifulSoup(r.text, "html.parser")
        enumerations = _parse_cf28_enumerations(soup)
        if enumerations:
            print(f"[{_timestamp()}] Enumerações cf_28 carregadas: {len(enumerations)}")
            return enumerations

        print(f"[{_timestamp()}] WARN: nenhuma enumeração cf_28 encontrada na página; usando fallback.")
    except Exception as e:
        print(f"[{_timestamp()}] WARN: falha ao buscar enumerações cf_28: {e}")

    return _default_cf28_enumerations()


def fetch_cf28_options(session: requests.Session) -> List[Tuple[str, str]]:
    """Backward-compatible helper returning (label, id) pairs."""
    return [(e["label"], e["id"]) for e in fetch_cf28_enumerations(session)]


def _absolute_url(href: str) -> str:
    if not href:
        return ""
    if href.startswith("http"):
        return href
    if href.startswith("/"):
        parsed = urlparse(PGE_BASE_URL)
        return f"{parsed.scheme}://{parsed.netloc}{href}"
    return urljoin(PGE_BASE_URL.rstrip("/") + "/", href)


def _issue_url(issue_id: str) -> str:
    return _absolute_url(f"{PGE_BASE_PATH}/issues/{issue_id}")


def _extract_attribute_value(soup: BeautifulSoup, *label_keywords: str) -> str:
    keywords = [k.lower() for k in label_keywords]
    for attr in soup.select(".attributes .attribute"):
        label = attr.select_one(".label")
        if not label:
            continue
        label_text = label.get_text(strip=True).rstrip(":").lower()
        if any(kw in label_text for kw in keywords):
            val = attr.select_one(".value")
            if val:
                text = val.get_text(strip=True)
                if text:
                    return text
    return ""


def _extract_catalogo(soup: BeautifulSoup) -> str:
    text = _extract_attribute_value(soup, "catálogo", "catalogo")
    return text or "sem_catalogo"


def _extract_titulo(soup: BeautifulSoup) -> str:
    h3 = soup.select_one("div.subject h3") or soup.select_one("#issue div.subject h3")
    return h3.get_text(strip=True) if h3 else ""


def _extract_ust(soup: BeautifulSoup) -> str:
    raw = _extract_attribute_value(soup, "ust")
    return _parse_decimal(raw)


def _extract_perfil(soup: BeautifulSoup) -> str:
    return _extract_attribute_value(soup, "perfil")


def _extract_assigned_user(soup: BeautifulSoup) -> str:
    user_link = (
        soup.select_one(".assigned-to .value a.user")
        or soup.select_one(".attribute.assigned-to a.user")
        or soup.select_one("div.attribute.assigned-to .value a")
    )
    if user_link:
        name = user_link.get_text(strip=True)
        if name:
            return name
    return "sem_usuario"


def _extract_pdf_url(soup: BeautifulSoup, issue_id: str) -> str:
    pdf_link = soup.select_one("a.pdf[href]")
    if pdf_link and pdf_link.get("href"):
        return _absolute_url(pdf_link["href"])
    return _absolute_url(f"{PGE_BASE_PATH}/issues/{issue_id}.pdf")


def _extract_commit_url(soup: BeautifulSoup) -> str:
    for journal in soup.select("#history div.journal"):
        notes_div = journal.select_one(".wiki.notes") or journal.select_one(".wiki")
        if not notes_div:
            continue

        for a in notes_div.find_all("a", href=True):
            href = a.get("href", "").strip()
            if not href.startswith("http"):
                continue
            prev_text = ""
            for prev in a.previous_siblings:
                if isinstance(prev, str):
                    prev_text = prev + prev_text
                else:
                    prev_text = prev.get_text() + prev_text
            if re.search(r"Commit\s*:", prev_text, re.IGNORECASE):
                return href

        text = notes_div.get_text("\n")
        m = re.search(r"Commit\s*:\s*(https?://\S+)", text, re.IGNORECASE)
        if m:
            return m.group(1).rstrip(".,;)")

    return ""


def _find_arquivos_table(soup: BeautifulSoup):
    for strong in soup.find_all("strong"):
        if "arquivos" in strong.get_text(strip=True).lower():
            parent = strong.find_parent()
            if not parent:
                continue
            table = parent.find("table")
            if table:
                return table
            table = parent.find_next("table")
            if table:
                return table
    return None


def _attachment_filename(link, download_href: str) -> str:
    label = link.select_one("span.icon-label")
    if label:
        name = label.get_text(strip=True)
        if name:
            return name
    basename = os.path.basename(urllib.parse.urlparse(download_href).path)
    return basename or "arquivo"


def _extract_arquivos_attachments(soup: BeautifulSoup) -> List[Tuple[str, str]]:
    table = _find_arquivos_table(soup)
    if not table:
        return []

    seen_urls = set()
    attachments: List[Tuple[str, str]] = []

    for tr in table.find_all("tr"):
        download_link = tr.select_one("a.icon-download[href]")
        if download_link and download_link.get("href"):
            href = download_link["href"]
            url = _absolute_url(href)
            if url in seen_urls:
                continue
            seen_urls.add(url)
            attachments.append((url, _attachment_filename(download_link, href)))
            continue

        attachment_link = tr.select_one("a.icon-attachment[href]")
        if attachment_link and attachment_link.get("href"):
            href = attachment_link["href"]
            download_href = href.replace("/attachments/", "/attachments/download/", 1)
            url = _absolute_url(download_href)
            if url in seen_urls:
                continue
            seen_urls.add(url)
            attachments.append((url, _attachment_filename(attachment_link, download_href)))

    return attachments


def _download_issue_attachments(
    session: requests.Session,
    soup: BeautifulSoup,
    issue_dir: str,
    issue_id: str,
) -> int:
    attachments = _extract_arquivos_attachments(soup)
    if not attachments:
        return 0

    downloaded = 0
    for url, name in attachments:
        if not _app_run:
            break
        dest = os.path.join(issue_dir, _valid_filename(name))
        _download_file(session, url, dest)
        if os.path.isfile(dest) and os.path.getsize(dest) > 0:
            downloaded += 1

    if downloaded:
        print(f"[{_timestamp()}] Issue {issue_id}: {downloaded} arquivo(s) baixado(s)")
    return downloaded


def _extract_issue_ids_from_list(soup: BeautifulSoup) -> List[str]:
    ids = []
    for tr in soup.select("table.list.issues tbody tr[id^='issue-']"):
        tr_id = tr.get("id", "")
        if tr_id.startswith("issue-"):
            issue_id = tr_id.replace("issue-", "").strip()
            if issue_id.isdigit():
                ids.append(issue_id)
                continue
        td_id = tr.find("td", class_="id")
        if td_id:
            link = td_id.find("a", href=re.compile(r"/issues/\d+"))
            if link:
                m = re.search(r"/issues/(\d+)", link.get("href", ""))
                if m:
                    ids.append(m.group(1))
    return ids


def _find_list_next_page(soup: BeautifulSoup) -> Optional[str]:
    next_li = soup.select_one("span.pagination li.next.page a")
    if next_li and next_li.get("href"):
        return _absolute_url(next_li["href"])
    for a in soup.select("span.pagination a"):
        text = a.get_text(strip=True)
        if "Próximo" in text and a.get("href"):
            return _absolute_url(a["href"])
    return None


def collect_all_issue_ids(session: requests.Session, filter_url: str) -> List[str]:
    """Paginates the filtered list via 'Próximo »' and collects all issue IDs."""
    all_ids: List[str] = []
    url: Optional[str] = filter_url
    page_num = 1

    while url and _app_run:
        print(f"[{_timestamp()}] Fetching list page {page_num}: {url}")
        r = _http_get(session, url)
        soup = BeautifulSoup(r.text, "html.parser")
        page_ids = _extract_issue_ids_from_list(soup)

        if not page_ids:
            print(f"[{_timestamp()}] Page {page_num} has no issues. Stopping list pagination.")
            break

        all_ids.extend(page_ids)
        print(f"[{_timestamp()}] Page {page_num}: {len(page_ids)} issues (total so far: {len(all_ids)})")

        next_url = _find_list_next_page(soup)
        if not next_url or next_url == url:
            break
        url = next_url
        page_num += 1

    return all_ids


def _extract_issue_id_from_page(soup: BeautifulSoup) -> Optional[str]:
    heading = soup.select_one("div#issue h2") or soup.find("h2")
    if heading:
        m = re.search(r"#(\d+)", heading.get_text())
        if m:
            return m.group(1)
    body = soup.find("body")
    if body and body.get("class"):
        for cls in body.get("class", []):
            if cls.startswith("issue-"):
                return cls.replace("issue-", "")
    return None


def _find_next_issue_url(soup: BeautifulSoup) -> Optional[str]:
    for a in soup.select("a[accesskey='n']"):
        href = a.get("href", "")
        if href and "/issues/" in href:
            return _absolute_url(href)

    for a in soup.find_all("a", href=True):
        text = a.get_text(strip=True)
        if re.search(r"Próxim", text, re.IGNORECASE) and "/issues/" in a["href"]:
            return _absolute_url(a["href"])

    for a in soup.find_all("a", rel="next"):
        href = a.get("href", "")
        if href and "/issues/" in href:
            return _absolute_url(href)

    return None


def process_issue_page(
    session: requests.Session,
    soup: BeautifulSoup,
    issue_id: str,
    evidencias_dir: str,
    individual_pdfs_dir: str,
) -> dict:
    catalogo = _extract_catalogo(soup)
    usuario = _extract_assigned_user(soup)
    titulo = _extract_titulo(soup)
    ust = _extract_ust(soup)
    perfil = _extract_perfil(soup)
    commit_url = _extract_commit_url(soup)
    pdf_url = _extract_pdf_url(soup, issue_id)
    issue_url = _issue_url(issue_id)

    issue_dir = os.path.join(
        evidencias_dir,
        _safe_dirname(catalogo),
        _safe_dirname(usuario),
        issue_id,
    )
    _safe_mkdir(issue_dir)

    pdf_dest = os.path.join(issue_dir, f"{issue_id}.pdf")
    _download_file(session, pdf_url, pdf_dest)
    if os.path.isfile(pdf_dest) and os.path.getsize(pdf_dest) > 0:
        flat_dest = os.path.join(individual_pdfs_dir, f"{issue_id}.pdf")
        shutil.copy2(pdf_dest, flat_dest)
    _download_issue_attachments(session, soup, issue_dir, issue_id)

    print(
        f"[{_timestamp()}] Issue {issue_id} | {catalogo} | {usuario}"
        + (f" | Commit: {commit_url}" if commit_url else "")
    )

    return {
        "ID": issue_id,
        "URL": issue_url,
        "Commit": commit_url,
        "Catalogo": catalogo,
        "Usuario": usuario,
        "Titulo": titulo,
        "UST": ust,
        "Perfil": perfil,
    }


def _download_filter_table_pdf(
    session: requests.Session,
    cf_28_id: str,
    filter_url: str,
    dest_path: str,
) -> None:
    pdf_url = build_pge_filter_pdf_url(cf_28_id)
    print(f"\n[{_timestamp()}] Gerando PDF da consulta...")
    _download_file(session, pdf_url, dest_path)

    if os.path.isfile(dest_path) and os.path.getsize(dest_path) > 0:
        return

    print(f"[{_timestamp()}] WARN: PDF via URL construída falhou; tentando scrape da lista.")
    try:
        r = _http_get(session, filter_url)
        soup = BeautifulSoup(r.text, "html.parser")
        pdf_link = soup.select_one("a.pdf[href]")
        if pdf_link and pdf_link.get("href"):
            scraped_url = _absolute_url(pdf_link["href"])
            _download_file(session, scraped_url, dest_path)
    except Exception as e:
        print(f"[{_timestamp()}] ERROR: fallback PDF scrape failed: {e}")


def _visit_issue_direct(session: requests.Session, issue_id: str) -> Tuple[Optional[BeautifulSoup], str]:
    url = _issue_url(issue_id)
    r = _http_get(session, url)
    soup = BeautifulSoup(r.text, "html.parser")
    return soup, url


def iterate_issues_via_proxima(
    session: requests.Session,
    first_issue_id: str,
    all_ids: List[str],
    evidencias_dir: str,
    individual_pdfs_dir: str,
) -> List[dict]:
    """Opens first issue and chains 'Próxima'; falls back to direct visits for missed IDs."""
    visited: List[dict] = []
    visited_set = set()

    url: Optional[str] = _issue_url(first_issue_id)
    chained = True

    while url and _app_run:
        print(f"[{_timestamp()}] Opening issue: {url}")
        r = _http_get(session, url)
        soup = BeautifulSoup(r.text, "html.parser")
        issue_id = _extract_issue_id_from_page(soup)
        if not issue_id:
            m = re.search(r"/issues/(\d+)", url)
            issue_id = m.group(1) if m else None
        if not issue_id:
            print(f"[{_timestamp()}] WARN: could not parse issue ID from {url}")
            break

        if issue_id not in visited_set:
            row = process_issue_page(session, soup, issue_id, evidencias_dir, individual_pdfs_dir)
            visited.append(row)
            visited_set.add(issue_id)

        next_url = _find_next_issue_url(soup)
        if next_url and next_url not in {_issue_url(v["ID"]) for v in visited}:
            url = next_url
        else:
            chained = False
            break

    if not chained or len(visited_set) < len(all_ids):
        missed = [i for i in all_ids if i not in visited_set]
        if missed:
            print(f"[{_timestamp()}] Fallback: visiting {len(missed)} issue(s) not reached via Próxima.")
        for issue_id in missed:
            if not _app_run:
                break
            soup, _issue_url_direct = _visit_issue_direct(session, issue_id)
            row = process_issue_page(session, soup, issue_id, evidencias_dir, individual_pdfs_dir)
            visited.append(row)
            visited_set.add(issue_id)

    return visited


def _process_issue_by_id(
    base_session: requests.Session,
    issue_id: str,
    evidencias_dir: str,
    individual_pdfs_dir: str,
) -> Optional[dict]:
    if not _app_run:
        return None
    try:
        session = _clone_session(base_session)
        soup, _ = _visit_issue_direct(session, issue_id)
        return process_issue_page(session, soup, issue_id, evidencias_dir, individual_pdfs_dir)
    except Exception as e:
        print(f"[{_timestamp()}] ERROR: issue {issue_id}: {e}")
        return None


def process_issues_parallel(
    session: requests.Session,
    all_ids: List[str],
    evidencias_dir: str,
    individual_pdfs_dir: str,
    max_workers: int,
) -> List[dict]:
    visited: List[dict] = []
    with ThreadPoolExecutor(max_workers=max_workers) as executor:
        futures = {
            executor.submit(
                _process_issue_by_id, session, issue_id, evidencias_dir, individual_pdfs_dir
            ): issue_id
            for issue_id in all_ids
        }
        for future in as_completed(futures):
            if not _app_run:
                for f in futures:
                    f.cancel()
                break
            row = future.result()
            if row:
                visited.append(row)
    return visited


def _build_tarefas_dataframe(visited: List[dict]) -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "ID": row["ID"],
                "Colaborador Atribuído": row["Usuario"],
                "Título": row["Titulo"],
                "UST": row["UST"],
                "Catálogo": row["Catalogo"],
                "Perfil": row["Perfil"],
            }
            for row in visited
        ],
        columns=_TAREFAS_COLUMNS,
    )


def _write_styled_excel(df: pd.DataFrame, path: str) -> None:
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    df.to_excel(path, index=False, engine="openpyxl")
    wb = load_workbook(path)
    ws = wb.active

    header_fill = PatternFill(start_color="AFC6E9", end_color="AFC6E9", fill_type="solid")
    row_fill_even = PatternFill(start_color="DEEAF6", end_color="DEEAF6", fill_type="solid")
    row_fill_odd = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    header_font = Font(bold=True, size=12)
    data_font = Font(size=11)

    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_idx in range(2, ws.max_row + 1):
        fill = row_fill_even if row_idx % 2 == 0 else row_fill_odd
        for cell in ws[row_idx]:
            cell.font = data_font
            cell.fill = fill
            cell.alignment = Alignment(vertical="center", wrap_text=True)

    for col in ws.columns:
        max_length = 0
        col_letter = col[0].column_letter
        for cell in col:
            if cell.value is not None:
                max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_length + 2, 60)

    wb.save(path)


def generate_evidences_pge(session, cf_28_id: str, month_label: str) -> None:
    global _app_run
    _app_run = True

    if not session:
        print(f"[{_timestamp()}] ERROR: Sessão inválida. Faça login primeiro.")
        return

    folder_stamp = datetime.date.today().strftime("%d-%m") + "-" + time.strftime("%H_%M")
    safe_month = re.sub(r'[<>:"/\\|?*]', "_", month_label)[:80]
    base_out_dir = os.path.join(".", "relatorios_pge", safe_month, folder_stamp)
    _safe_mkdir(base_out_dir)
    evidencias_dir = os.path.join(base_out_dir, EVIDENCIAS_DIR)
    _safe_mkdir(evidencias_dir)
    individual_pdfs_dir = os.path.join(base_out_dir, RELATORIO_INDIVIDUAIS_DIR)
    _safe_mkdir(individual_pdfs_dir)

    print(f"\n[{_timestamp()}] Preparando evidências PGE...")
    print(f"Entregue em: {month_label} (cf_28={cf_28_id})")
    print(f"Pasta de saída: {base_out_dir}")

    filter_url = build_pge_filter_url(cf_28_id)
    print(f"\n[{_timestamp()}] URL do filtro: {filter_url}")

    all_ids = collect_all_issue_ids(session, filter_url)
    if not all_ids:
        print(f"[{_timestamp()}] Nenhuma issue encontrada. Encerrando.")
        return

    print(f"\n[{_timestamp()}] Total de issues na lista: {len(all_ids)}")

    workers = _issue_workers()
    print(f"[{_timestamp()}] Processando issues com {workers} worker(s)")

    first_issue_id = all_ids[0]
    if workers == 1:
        visited = iterate_issues_via_proxima(
            session, first_issue_id, all_ids, evidencias_dir, individual_pdfs_dir
        )
    else:
        print(f"[{_timestamp()}] Modo paralelo: {workers} workers, {len(all_ids)} issues")
        visited = process_issues_parallel(
            session, all_ids, evidencias_dir, individual_pdfs_dir, workers
        )

    if visited:
        df_tarefas = _build_tarefas_dataframe(visited)

        csv_path = os.path.join(base_out_dir, TAREFAS_PROCESSADAS_NOME)
        df_tarefas.to_csv(csv_path, sep=";", index=False, encoding="utf-8-sig")
        print(f"\n[{_timestamp()}] Tarefas processadas: {len(visited)}")
        print(f"[{_timestamp()}] CSV salvo: {csv_path}")

        if _app_run:
            relatorio_pdf_path = os.path.join(base_out_dir, RELATORIO_PDF_NOME)
            _download_filter_table_pdf(session, cf_28_id, filter_url, relatorio_pdf_path)

        if _app_run:
            xlsx_path = os.path.join(base_out_dir, TAREFAS_GERAL_NOME)
            try:
                _write_styled_excel(df_tarefas, xlsx_path)
                print(f"[{_timestamp()}] Excel salvo: {xlsx_path}")
            except ImportError:
                csv_fallback = os.path.join(base_out_dir, "TAREFAS_GERAL.csv")
                df_tarefas.to_csv(csv_fallback, sep=";", index=False, encoding="utf-8-sig")
                print(
                    f"[{_timestamp()}] WARN: openpyxl não instalado; "
                    f"Excel indisponível. CSV salvo: {csv_fallback}"
                )
                print(f"[{_timestamp()}] Instale com: pip install openpyxl")

    print(f"\n[{_timestamp()}] Fim! DIRETORIO_FINAL:{base_out_dir}")
